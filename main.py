import os, argparse
from openpyxl import Workbook

# from tables.comm_properties import export_commproperties, SHEET_NAME as COMM_SHEET
# from tables.lock_outs import export_lockouts, SHEET_NAME as LOCK_OUTS_SHEET
# from tables.res_unit_amenities import export_resunitamenities, SHEET_NAME as RES_UNIT_AMENITIES_SHEET
from structure.printer import save_nested_structure_as_json
from tables.export_all import export_all_tables


INPUT_DIR = "input"
OUTPUT_DIR = "output"

def main():
    p = argparse.ArgumentParser(description="XML → Excel (one sheet per table)")
    p.add_argument("filename", help="XML filename only (in input/)")
    p.add_argument("--print-structure", action="store_true",
                   help="Print nested XML structure and exit")
    p.add_argument("--max-depth", type=int, default=1,
                   help="Max depth for --print-structure")
    p.add_argument("--max-per-tag-per-level", type=int, default=2,
                   help="Max occurrences per tag per level")

    args = p.parse_args()

    xml_path = os.path.join(INPUT_DIR, args.filename)
    if not os.path.isfile(xml_path):
        raise SystemExit(f"❌ Input file not found: {xml_path}")

    if args.print_structure:
        save_nested_structure_as_json(
            xml_path,
            max_depth=args.max_depth,
            max_per_tag_per_level=args.max_per_tag_per_level,
        )

        return

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    out_path = os.path.join(
        OUTPUT_DIR, os.path.splitext(args.filename)[0] + ".xlsx"
    )

    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    # ws = wb.create_sheet(COMM_SHEET)
    # export_commproperties(xml_path, ws)

    # ws = wb.create_sheet(LOCK_OUTS_SHEET)
    # export_lockouts(xml_path, ws)

    # ws = wb.create_sheet(RES_UNIT_AMENITIES_SHEET)
    # export_resunitamenities(xml_path, ws)
    export_all_tables(xml_path, wb)

    wb.save(out_path)
    print(f"✅ Excel written to: {out_path}")

if __name__ == "__main__":
    main()
